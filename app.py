import json
import os
import re
from typing import Any

import streamlit as st
from google import genai
from pypdf import PdfReader
from docx import Document

APP_TITLE = "Resume ATS Checker"
DEFAULT_MODEL = "gemini-2.5-flash"


def get_saved_api_key() -> str:
    """Read a key from Streamlit Secrets or environment, without exposing it."""
    try:
        secret_key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        secret_key = ""
    return str(secret_key or os.getenv("GEMINI_API_KEY", ""))
MAX_FILE_BYTES = 5 * 1024 * 1024
SUPPORTED_TYPES = ["pdf", "docx", "txt"]

st.set_page_config(page_title=APP_TITLE, page_icon="📄", layout="wide")


def extract_resume_text(uploaded_file) -> str:
    """Extract readable text from a PDF, DOCX, or TXT upload."""
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()

    if len(raw) > MAX_FILE_BYTES:
        raise ValueError("The file is larger than 5 MB. Please upload a smaller file.")

    if name.endswith(".pdf"):
        from io import BytesIO
        reader = PdfReader(BytesIO(raw))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        from io import BytesIO
        doc = Document(BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = raw.decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 80:
        raise ValueError(
            "I couldn't extract enough text from this file. If it is a scanned PDF, "
            "run OCR on it first, or upload a text-based PDF/DOCX/TXT file."
        )
    if len(text) > 30000:
        text = text[:30000]
    return text


def parse_json_response(text: str) -> dict[str, Any]:
    """Parse JSON, tolerating a surrounding Markdown code fence."""
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    data = json.loads(cleaned)
    if not isinstance(data, dict):
        raise ValueError("The AI response was not a JSON object.")
    return data


def analyze_resume(resume_text: str, job_description: str, api_key: str, model: str) -> dict[str, Any]:
    client = genai.Client(api_key=api_key)
    prompt = f"""
You are a careful resume reviewer. Evaluate the resume for ATS-readability and job relevance.
This is an estimate, not a score from any real Applicant Tracking System. Do not infer facts
that are not present. Never invent experience, qualifications, metrics, or skills.

Return ONLY valid JSON with this exact shape:
{{
  "ats_score": 0,
  "score_explanation": "brief explanation of why this estimated score was assigned",
  "summary": "2-4 sentence overall review",
  "strengths": ["specific strength", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "issue": "specific issue", "recommendation": "actionable fix"}}
  ],
  "formatting_checks": [
    {{"item": "check name", "status": "Good|Needs attention|Not clear", "notes": "brief note"}}
  ],
  "keyword_gaps": ["relevant missing keyword or phrase, if supported by job description"],
  "rewritten_summary": "a truthful, improved professional summary based only on the resume; if details are insufficient, provide a clearly marked template with placeholders",
  "bullet_rewrites": [
    {{"before": "short original bullet or description", "after": "improved truthful version", "note": "what changed"}}
  ]
}}

Rules:
- ats_score must be an integer from 0 to 100.
- Be constructive, specific, and concise.
- Assess readability, section clarity, skills, evidence of impact, consistency, and alignment to the job description when provided.
- Do not penalize a candidate for personal characteristics or protected traits.
- Do not claim to know the exact scoring rules of any ATS vendor.
- If no job description is supplied, assess general ATS readiness and set keyword_gaps to [].
- Keep the output useful to a student or early-career applicant when relevant.
- Limit strengths to 6, improvements to 8, formatting_checks to 8, keyword_gaps to 12, and bullet_rewrites to 4.
- In rewrites, preserve the original meaning and do not fabricate numbers or achievements.

JOB DESCRIPTION (may be empty):
{job_description[:10000] if job_description.strip() else "[Not provided]"}

RESUME TEXT:
{resume_text}
"""
    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config={"response_mime_type": "application/json"},
    )
    if not getattr(response, "text", None):
        raise RuntimeError("Gemini returned an empty response. Please try again.")
    result = parse_json_response(response.text)
    score = result.get("ats_score")
    if not isinstance(score, (int, float)) or not 0 <= score <= 100:
        raise ValueError("The AI returned an invalid score. Please try again.")
    result["ats_score"] = int(score)
    for key in ("strengths", "improvements", "formatting_checks", "keyword_gaps", "bullet_rewrites"):
        if not isinstance(result.get(key, []), list):
            result[key] = []
    return result


def safe_text(value: Any, fallback: str = "Not provided") -> str:
    if value is None:
        return fallback
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


st.title("📄 Resume ATS Checker")
st.caption("Upload a resume to get an estimated ATS-readiness score and practical improvement suggestions.")
st.info(
    "Privacy note: the extracted resume text is sent to Google Gemini to generate the review. "
    "Avoid uploading information you do not want processed. This app does not intentionally save "
    "your uploaded file or analysis to a database."
)

with st.sidebar:
    st.header("Settings")
    api_key = st.text_input(
        "Gemini API key",
        type="password",
        value=get_saved_api_key(),
        help="For local use, paste your key here. For deployment, add GEMINI_API_KEY in Streamlit Secrets.",
    )
    model = st.text_input("Gemini model", value=DEFAULT_MODEL)
    st.caption("Never commit API keys to GitHub.")

uploaded_file = st.file_uploader(
    "Choose your resume",
    type=SUPPORTED_TYPES,
    help="Supported formats: PDF, DOCX, TXT. Maximum size: 5 MB.",
)
job_description = st.text_area(
    "Target job description (optional)",
    height=180,
    placeholder="Paste the job description here to get a more role-specific review.",
)

if uploaded_file:
    st.write(f"**Selected file:** {uploaded_file.name} · {uploaded_file.size / 1024:.1f} KB")

if st.button("Analyze resume", type="primary", disabled=uploaded_file is None):
    if not api_key.strip():
        st.error("Add your Gemini API key in the sidebar first.")
    elif not model.strip():
        st.error("Enter a Gemini model name in the sidebar.")
    else:
        try:
            with st.spinner("Extracting resume text and asking Gemini to review it..."):
                resume_text = extract_resume_text(uploaded_file)
                result = analyze_resume(resume_text, job_description, api_key.strip(), model.strip())
            st.session_state["resume_analysis"] = result
            st.session_state["resume_filename"] = uploaded_file.name
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            st.caption("Check your API key, model name, internet connection, file format, and Gemini quota.")

result = st.session_state.get("resume_analysis")
if result:
    st.divider()
    st.subheader(f"Results for {st.session_state.get('resume_filename', 'your resume')}")
    score_col, explanation_col = st.columns([1, 3])
    with score_col:
        st.metric("Estimated ATS score", f"{result.get('ats_score', 0)}/100")
        st.progress(max(0, min(100, int(result.get("ats_score", 0)))) / 100)
    with explanation_col:
        st.write(safe_text(result.get("score_explanation")))
        st.write(safe_text(result.get("summary")))

    tab1, tab2, tab3, tab4 = st.tabs(
        ["Improvements", "Strengths & keywords", "Formatting checks", "Rewrite suggestions"]
    )
    with tab1:
        improvements = result.get("improvements", [])
        if improvements:
            for item in improvements:
                if not isinstance(item, dict):
                    st.write(f"• {safe_text(item)}")
                    continue
                priority = safe_text(item.get("priority"), "Medium")
                st.markdown(f"**{priority} priority: {safe_text(item.get('issue'))}**")
                st.write(safe_text(item.get("recommendation")))
                st.divider()
        else:
            st.write("No improvement suggestions were returned.")
    with tab2:
        st.markdown("**Strengths**")
        strengths = result.get("strengths", [])
        if strengths:
            for item in strengths:
                st.write(f"• {safe_text(item)}")
        else:
            st.write("No strengths were returned.")
        st.markdown("**Potential keyword gaps**")
        gaps = result.get("keyword_gaps", [])
        if gaps:
            for item in gaps:
                st.write(f"• {safe_text(item)}")
        else:
            st.write("No specific keyword gaps identified, or no job description was provided.")
    with tab3:
        for item in result.get("formatting_checks", []):
            if isinstance(item, dict):
                st.markdown(f"**{safe_text(item.get('item'))}: {safe_text(item.get('status'))}**")
                st.write(safe_text(item.get("notes")))
            else:
                st.write(f"• {safe_text(item)}")
    with tab4:
        st.markdown("**Suggested professional summary**")
        st.write(safe_text(result.get("rewritten_summary")))
        st.markdown("**Bullet point rewrites**")
        rewrites = result.get("bullet_rewrites", [])
        if rewrites:
            for item in rewrites:
                if isinstance(item, dict):
                    st.markdown(f"**Before:** {safe_text(item.get('before'))}")
                    st.markdown(f"**After:** {safe_text(item.get('after'))}")
                    st.caption(safe_text(item.get("note"), ""))
                    st.divider()
        else:
            st.write("No bullet rewrites were returned.")

    downloadable = json.dumps(result, ensure_ascii=False, indent=2)
    st.download_button(
        "Download analysis as JSON",
        data=downloadable,
        file_name="resume_analysis.json",
        mime="application/json",
    )
    st.caption(
        "This score is an AI-generated estimate, not a guarantee of passing an ATS or getting an interview. "
        "Review every suggestion and keep all statements accurate to your experience."
    )
else:
    st.markdown("### How it works")
    st.write("1. Upload a text-based PDF, DOCX, or TXT resume.")
    st.write("2. Optionally paste a job description for targeted keyword feedback.")
    st.write("3. Review the estimated score, strengths, issues, and rewrite suggestions.")
